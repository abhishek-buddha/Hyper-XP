import re
from datetime import datetime
from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_YELLOW   = PatternFill("solid", fgColor="FFFF00")
_RED      = PatternFill("solid", fgColor="FF6B6B")
_ROW_FAIL = PatternFill("solid", fgColor="FFDEDE")
_ROW_WARN = PatternFill("solid", fgColor="FFF3CD")
_BOLD     = Font(bold=True)

_GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")
_RED_FILL   = PatternFill("solid", fgColor="FFDEDE")
_AMBER_FILL = PatternFill("solid", fgColor="FFF3CD")
_GRAY_FILL  = PatternFill("solid", fgColor="F0F0F0")


def _auto_fit_columns(ws) -> None:
    for col in ws.columns:
        if not col:
            continue
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def generate_workbook(sheets: List[dict]) -> bytes:
    """Create an Excel workbook with one sheet per entry in sheets[].

    Each sheet entry: { name, columns: [...], rows: [{col: {value, confidence}, _row_validation: {status, reason}}] }
    Row fill priority: fail > warning > cell-level (null=red, low confidence=yellow).
    _row_validation key is never written as a column.
    """
    if not sheets:
        raise ValueError("sheets must contain at least one entry")
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_def in sheets:
        ws = wb.create_sheet(title=sheet_def["name"][:31])
        columns = sheet_def.get("columns", [])

        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = _BOLD
        ws.freeze_panes = "A2"

        for row_idx, row_data in enumerate(sheet_def.get("rows", []), 2):
            rv = row_data.get("_row_validation", {})
            row_fill = None
            if rv.get("status") == "fail":
                row_fill = _ROW_FAIL
            elif rv.get("status") == "warning":
                row_fill = _ROW_WARN

            for col_idx, col_name in enumerate(columns, 1):
                cell_obj   = row_data.get(col_name, {})
                value      = cell_obj.get("value")
                confidence = cell_obj.get("confidence", "high")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if row_fill:
                    cell.fill = row_fill
                elif value is None:
                    cell.fill = _RED
                elif confidence == "low":
                    cell.fill = _YELLOW

        _auto_fit_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_summary_sheet(ws, document_type: str, sheets: list, validation_summary: dict) -> None:
    """Write a QC summary onto an existing worksheet."""
    for col, width in zip("ABCDEF", [26, 16, 16, 16, 16, 16]):
        ws.column_dimensions[col].width = width

    def _hdr(row, col, value, fill=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(bold=True, size=11)
        if fill:
            c.fill = fill
        return c

    def _val(row, col, value, fill=None):
        c = ws.cell(row=row, column=col, value=value)
        if fill:
            c.fill = fill
        return c

    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = "HYPERXP — QC VALIDATION REPORT"
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(horizontal="center")

    ws["A3"] = "Document"
    ws["A3"].font = _BOLD
    ws["B3"] = document_type
    ws.merge_cells("B3:F3")

    ws["A4"] = "Generated"
    ws["A4"].font = _BOLD
    ws["B4"] = datetime.now().strftime("%d/%m/%Y  %H:%M")

    _hdr(6, 1, "OVERALL SUMMARY")

    s = validation_summary or {}
    ops_sheets = [sh for sh in sheets if not re.search(r"shift|signator|parameter record", sh.get("name", ""), re.I)]
    total_rows = sum(len(sh.get("rows", [])) for sh in ops_sheets)
    na_rows = total_rows - (s.get("total") or 0)

    for i, (label, val, fill) in enumerate([
        ("Total rows reviewed", total_rows,          None),
        ("Passed",              s.get("passed",   0), _GREEN_FILL),
        ("Failed",              s.get("failed",   0), _RED_FILL),
        ("Warnings",            s.get("warnings", 0), _AMBER_FILL),
        ("N/A (admin steps)",   na_rows,               _GRAY_FILL),
    ], start=7):
        ws.cell(row=i, column=1, value=label).font = Font(bold=(i == 7))
        c = ws.cell(row=i, column=2, value=val)
        if fill:
            ws.cell(row=i, column=1).fill = fill
            c.fill = fill

    br = 14
    _hdr(br, 1, "SHEET BREAKDOWN")
    br += 1
    for col, label in enumerate(["Sheet", "Passed", "Failed", "Warnings", "N/A", "Total"], 1):
        _hdr(br, col, label)
    br += 1

    for sh in ops_sheets:
        counts = {"pass": 0, "fail": 0, "warning": 0, "na": 0}
        for row in sh.get("rows", []):
            st = row.get("_row_validation", {}).get("status", "na")
            counts[st] = counts.get(st, 0) + 1
        total = sum(counts.values())
        _val(br, 1, sh["name"])
        _val(br, 2, counts["pass"],    _GREEN_FILL if counts["pass"]    else None)
        _val(br, 3, counts["fail"],    _RED_FILL   if counts["fail"]    else None)
        _val(br, 4, counts["warning"], _AMBER_FILL if counts["warning"] else None)
        _val(br, 5, counts["na"])
        _val(br, 6, total)
        br += 1


def generate_report(sheets: List[dict], document_type: str, validation_summary: dict) -> bytes:
    """Generate a single QC Summary Report Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "QC Report"
    _write_summary_sheet(ws, document_type, sheets, validation_summary)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_bulk_report(records: List[dict]) -> bytes:
    """Combine multiple QC reports into one workbook.

    records: [{document_type, sheets, validation_summary, batch_no, filename}]
    Sheet 1  — Master Summary (one row per record)
    Sheets 2+ — individual QC summary per record
    """
    wb = Workbook()

    # ── Master Summary sheet ───────────────────────────────────────────────
    master = wb.active
    master.title = "Master Summary"
    for col, width in zip("ABCDEFG", [30, 14, 10, 10, 12, 6, 20]):
        master.column_dimensions[col].width = width

    master.merge_cells("A1:G1")
    t = master["A1"]
    t.value = "HYPERXP — BULK QC REPORT"
    t.font = Font(bold=True, size=14)
    t.alignment = Alignment(horizontal="center")

    master["A3"] = "Generated"
    master["A3"].font = _BOLD
    master["B3"] = datetime.now().strftime("%d/%m/%Y  %H:%M")

    hdr_row = 5
    for col, label in enumerate(["Document", "Total", "Passed", "Failed", "Warnings", "N/A", "Batch / File"], 1):
        c = master.cell(row=hdr_row, column=col, value=label)
        c.font = Font(bold=True)

    for r, rec in enumerate(records, start=hdr_row + 1):
        s = rec.get("validation_summary") or {}
        ops = [sh for sh in rec.get("sheets", []) if not re.search(r"shift|signator|parameter record", sh.get("name", ""), re.I)]
        total = sum(len(sh.get("rows", [])) for sh in ops)
        na = total - (s.get("total") or 0)
        label = rec.get("batch_no") or rec.get("filename") or "—"

        master.cell(row=r, column=1, value=rec.get("document_type", ""))
        master.cell(row=r, column=2, value=total)
        c_pass = master.cell(row=r, column=3, value=s.get("passed", 0))
        c_fail = master.cell(row=r, column=4, value=s.get("failed", 0))
        c_warn = master.cell(row=r, column=5, value=s.get("warnings", 0))
        master.cell(row=r, column=6, value=na)
        master.cell(row=r, column=7, value=label)

        if s.get("passed", 0):  c_pass.fill = _GREEN_FILL
        if s.get("failed", 0):  c_fail.fill = _RED_FILL
        if s.get("warnings", 0): c_warn.fill = _AMBER_FILL

    # ── Individual report sheets ───────────────────────────────────────────
    seen: dict = {}
    for rec in records:
        raw = (rec.get("batch_no") or rec.get("document_type") or "Report")[:28]
        base = re.sub(r"[\\/*?:\[\]]", "_", raw)
        count = seen.get(base, 0) + 1
        seen[base] = count
        title = base if count == 1 else f"{base[:25]} ({count})"
        ws = wb.create_sheet(title=title)
        _write_summary_sheet(ws, rec.get("document_type", ""), rec.get("sheets", []), rec.get("validation_summary") or {})

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
